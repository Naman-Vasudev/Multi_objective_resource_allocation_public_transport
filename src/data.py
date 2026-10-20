"""Load, parse and validate CTU_Routes_Structured.xlsx.

Produces (in data/processed/):
    routes.csv          one row per route (68), the clean parameter table
    stops.csv           one row per stop (649) from the Stop Coverage sheet
    route_stops.csv     long table (route_no, stop) - which route serves which stop
    route_features.csv  per-route stop-network features

Rules:
    * The TOTAL row and blank/note rows are dropped; counts and sums are asserted.
    * route_no is ALWAYS a string ("1", "25/25A"), never an int.
    * Operating windows are read from the Operating Windows sheet (typed time cells
      and Excel's cached results). The Route Master text is parsed ONLY to
      cross-check, never as the source.
"""
from __future__ import annotations

import datetime as dt
import math
import re

import numpy as np
import openpyxl
import pandas as pd

from src import params as P

# ======================================================================
# Small pure parsers (unit-tested in tests/test_parsing.py)
# ======================================================================


def normalize_route_no(value) -> str:
    """Return the route number as a clean string. 1 -> "1", 1.0 -> "1", " 25/25A " -> "25/25A"."""
    if value is None:
        raise ValueError("route number is empty")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).strip()
    if not s:
        raise ValueError("route number is empty")
    return s


def parse_headway_range(text) -> tuple[float, float]:
    """Parse a published headway string [min] into (lo, hi) minutes.

    "20-25" -> (20.0, 25.0);  "15-30" -> (15.0, 30.0);  "20" -> (20.0, 20.0).
    lo is the SHORTEST headway (=> highest frequency f_max), hi the LONGEST (=> f_min).
    """
    s = str(text).strip().replace("–", "-").replace("—", "-")
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(?:-\s*(\d+(?:\.\d+)?))?", s)
    if not m:
        raise ValueError(f"cannot parse headway {text!r}")
    lo = float(m.group(1))
    hi = float(m.group(2)) if m.group(2) else lo
    if lo <= 0 or hi < lo:
        raise ValueError(f"bad headway range {text!r}")
    return lo, hi


_TIME = r"(\d{1,2})[.:](\d{2})"


def parse_operating_hours(text) -> list[tuple[dt.time, dt.time]]:
    """Parse Route Master operating-hours text into a list of (start, end) windows.

    "06.10-17.20 / 07.30-18.40" -> [(06:10, 17:20), (07:30, 18:40)]  (one window per direction)
    "10:10-15.50"               -> [(10:10, 15:50)]   (':' and '.' both accepted)
    "22.00-04.40"               -> [(22:00, 04:40)]   (end < start means past midnight)
    Trailing notes in parentheses, e.g. "(fixed departures)", are ignored.
    USED ONLY FOR CROSS-CHECKING the Operating Windows sheet.
    """
    s = re.sub(r"\(.*?\)", "", str(text)).strip()
    windows = []
    for part in s.split("/"):
        m = re.fullmatch(rf"\s*{_TIME}\s*-\s*{_TIME}\s*", part)
        if not m:
            raise ValueError(f"cannot parse operating window {part!r} in {text!r}")
        h1, m1, h2, m2 = map(int, m.groups())
        windows.append((dt.time(h1, m1), dt.time(h2, m2)))
    return windows


def parse_route_list(text) -> list[str]:
    """Split a Stop Coverage 'Route Numbers' cell on commas. "1, 25/25A, 2A" -> ["1", "25/25A", "2A"].

    The '/' inside "25/25A" is part of the route name and is NOT a separator.
    """
    if text is None or str(text).strip() == "":
        return []
    return [normalize_route_no(tok) for tok in str(text).split(",") if tok.strip()]


def _place_key(name: str) -> str:
    """Normalise a place name for equality tests ("Mani Majra" == "Manimajra")."""
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


def _minutes(t: dt.time) -> int:
    return t.hour * 60 + t.minute


def span_hours(first: dt.time, last: dt.time) -> float:
    """Daily service span [h]; if last < first the service runs past midnight (same rule as the sheet)."""
    d = _minutes(last) - _minutes(first)
    if d < 0:
        d += 24 * 60
    return d / P.MIN_PER_HOUR


# ======================================================================
# Workbook readers
# ======================================================================


def _open(path=P.RAW_XLSX):
    """Open twice: formulas (to know what is computed) and cached values (Excel's results)."""
    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    return wb_f, wb_v


FOOTER_LABELS = {"TOTAL", "ROUTES ACTIVE"}  # summary rows under the route tables


def _route_rows(ws, first_row: int) -> list[int]:
    """Row indices of real data rows: stop at the first blank cell or footer label (TOTAL, Routes active).

    Row counts are asserted by the callers, so an unexpected footer fails loudly rather than silently.
    """
    rows = []
    r = first_row
    while True:
        v = ws.cell(r, 1).value
        if v is None or str(v).strip() == "" or str(v).strip().upper() in FOOTER_LABELS:
            break
        rows.append(r)
        r += 1
    return rows


def read_route_master(wb_v) -> pd.DataFrame:
    """Route Master sheet -> DataFrame (68 rows, TOTAL row excluded)."""
    ws = wb_v["Route Master"]
    header = [c.value for c in ws[1]]
    assert header[0] == "Route No." and header[7] == "Route Length (km)", header
    recs = []
    for r in _route_rows(ws, 2):
        v = [ws.cell(r, c).value for c in range(1, 11)]
        recs.append(dict(
            route_no=normalize_route_no(v[0]), origin=str(v[1]).strip(), destination=str(v[2]).strip(),
            hours_text=str(v[3]).strip(), headway_text=str(v[4]).strip(),
            freq_min_col=float(v[5]), freq_max_col=float(v[6]),
            d_r_km=float(v[7]), x_base=int(v[8]), n_via_stops=int(v[9]),
        ))
    return pd.DataFrame(recs)


def read_operating_windows(wb_f, wb_v) -> pd.DataFrame:
    """Operating Windows sheet -> per-route windows, first/last departure, span [h], active flags.

    first/last/span come from Excel's cached formula results. If the file was re-saved by a tool
    that drops cached values, they are recomputed from the typed time cells (C, D, F, G) with the
    same rule as the sheet formulas (MIN of starts, MAX of ends, wrap past midnight).
    """
    ws_v, ws_f = wb_v["Operating Windows"], wb_f["Operating Windows"]
    hdr = [c.value for c in ws_v[6]]
    assert hdr[0] == "Route No." and hdr[9] == "Daily span (hrs)", hdr
    recs = []
    for r in _route_rows(ws_v, 7):
        g = lambda c: ws_v.cell(r, c).value  # noqa: E731
        a_first, a_last, b_first, b_last = g(3), g(4), g(6), g(7)
        starts = [t for t in (a_first, b_first) if t is not None]
        ends = [t for t in (a_last, b_last) if t is not None]
        first_re, last_re = min(starts, key=_minutes), max(ends, key=_minutes)
        first, last, span = g(8), g(9), g(10)
        cached = first is not None and last is not None and span is not None
        if not cached:
            first, last, span = first_re, last_re, span_hours(first_re, last_re)
        recs.append(dict(
            route_no=normalize_route_no(g(1)), dir_a=g(2), dir_b=g(5),
            a_first=a_first, a_last=a_last, b_first=b_first, b_last=b_last,
            first_departure=first, last_departure=last, span_hours=float(span),
            first_recomputed=first_re, last_recomputed=last_re,
            span_recomputed=span_hours(first_re, last_re), from_cache=cached,
            service_type=g(11),
            active_am=g(12), active_midday=g(13), active_pm=g(14), ow_note=g(15),
            span_formula=ws_f.cell(r, 10).value,
        ))
    return pd.DataFrame(recs)


def read_period_boundaries(wb_v) -> dict[str, tuple[float, float]]:
    """Operating Windows cells A2:C4 -> {period: (start_h, end_h)} in hours of the day [h].

    Raises DataError if they differ from params.PERIODS (the two must be kept in sync).
    """
    ws = wb_v["Operating Windows"]
    label_to_key = {v: k for k, v in P.PERIOD_SHEET_LABELS.items()}
    out = {}
    for r in range(2, 2 + len(P.PERIODS)):
        label, start, end = (ws.cell(r, c).value for c in (1, 2, 3))
        _require(label in label_to_key, f"Operating Windows A{r}: unexpected period label {label!r}")
        out[label_to_key[label]] = (_minutes(start) / P.MIN_PER_HOUR, _minutes(end) / P.MIN_PER_HOUR)
    _require(out == P.PERIODS,
             f"period boundaries in the sheet {out} differ from params.PERIODS {P.PERIODS}; "
             "edit one so they match (Operating Windows B2:C4 or src/params.py)")
    return out


def period_overlap_hours(first: dt.time, last: dt.time, start_h: float, end_h: float) -> float:
    """Hours [h] of the daily service window [first, last] that fall inside the period [start_h, end_h).

    A window with last < first runs past midnight; the period is also tested shifted by +24 h.
    """
    a = _minutes(first) / P.MIN_PER_HOUR
    b = _minutes(last) / P.MIN_PER_HOUR
    if b < a:
        b += 24.0
    return sum(max(0.0, min(b, e + k) - max(a, s + k)) for k in (0.0, 24.0) for s, e in [(start_h, end_h)])


def read_model_variables(wb_v) -> pd.DataFrame:
    """Model Variables sheet (cached values) -> DataFrame, blank/legend rows excluded."""
    ws = wb_v["Model Variables"]
    assert str(ws.cell(4, 1).value).startswith("Route"), ws.cell(4, 1).value
    recs = []
    for r in _route_rows(ws, 5):
        v = [ws.cell(r, c).value for c in range(1, 12)]
        recs.append(dict(route_no=normalize_route_no(v[0]), xl_d=v[1], xl_tau=v[2], xl_x=v[3],
                         xl_fmin=v[4], xl_fmax=v[5], xl_fmid=v[6], xl_xtau=v[7],
                         xl_D=v[8], xl_C=v[9], xl_o=v[10]))
    return pd.DataFrame(recs)


def read_pdf_verification(wb_v, rm: pd.DataFrame) -> pd.DataFrame:
    """PDF Verification sheet -> per-route match flags.

    Uses Excel's cached 'Match' column; if absent, re-does the same comparison
    (PDF frequency text, length, buses) against Route Master.
    """
    ws = wb_v["PDF Verification"]
    assert ws.cell(8, 1).value == "Route No.", ws.cell(8, 1).value
    rm_i = rm.set_index("route_no")
    recs = []
    for r in _route_rows(ws, 9):
        rn = normalize_route_no(ws.cell(r, 1).value)
        pdf_freq, pdf_len, pdf_bus = str(ws.cell(r, 2).value).strip(), ws.cell(r, 3).value, ws.cell(r, 4).value
        match = ws.cell(r, 8).value
        if match is None:  # no cached value -> recompute
            ok = (rn in rm_i.index and rm_i.at[rn, "headway_text"] == pdf_freq
                  and math.isclose(rm_i.at[rn, "d_r_km"], pdf_len) and rm_i.at[rn, "x_base"] == pdf_bus)
            match = "OK" if ok else "CHECK"
        recs.append(dict(route_no=rn, pdf_freq=pdf_freq, pdf_len=pdf_len, pdf_bus=pdf_bus, match=match))
    return pd.DataFrame(recs)


def read_stop_coverage(wb_v) -> pd.DataFrame:
    """Stop Coverage sheet -> one row per stop (blank row and trailing note excluded)."""
    ws = wb_v["Stop Coverage"]
    assert ws.cell(1, 1).value == "Stop / Landmark", ws.cell(1, 1).value
    recs = []
    for r in _route_rows(ws, 2):
        v = [ws.cell(r, c).value for c in range(1, 5)]
        recs.append(dict(stop=str(v[0]).strip(), n_routes_serving=int(v[1]),
                         buses_passing=int(v[2]), routes=parse_route_list(v[3])))
    return pd.DataFrame(recs)


# ======================================================================
# Build + validate
# ======================================================================


class DataError(AssertionError):
    """Raised when the workbook does not match the facts the model relies on."""


def _require(cond: bool, msg: str):
    if not cond:
        raise DataError(msg)


def build_routes(rm: pd.DataFrame, ow: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Merge Route Master + Operating Windows into the clean route table. Returns (routes, warnings)."""
    warnings: list[str] = []

    _require(len(rm) == P.N_ROUTES, f"Route Master has {len(rm)} routes, expected {P.N_ROUTES}")
    _require(rm.route_no.is_unique, "duplicate route numbers in Route Master")
    _require(int(rm.x_base.sum()) == P.FLEET_B, f"sum(x_base)={rm.x_base.sum()}, expected {P.FLEET_B}")
    _require(math.isclose(rm.d_r_km.sum(), P.TOTAL_LENGTH_KM, abs_tol=1e-6),
             f"sum(d_r)={rm.d_r_km.sum()}, expected {P.TOTAL_LENGTH_KM}")
    _require(list(ow.route_no) == list(rm.route_no) or set(ow.route_no) == set(rm.route_no),
             "Operating Windows routes differ from Route Master routes")

    # headways: numeric columns are the source; the text must agree with them
    for row in rm.itertuples():
        lo, hi = parse_headway_range(row.headway_text)
        _require((lo, hi) == (row.freq_min_col, row.freq_max_col),
                 f"route {row.route_no}: headway text {row.headway_text!r} != columns "
                 f"({row.freq_min_col}, {row.freq_max_col})")

    df = rm.merge(ow, on="route_no", how="left", validate="one_to_one")

    # operating windows: Route Master text must describe the same windows as the sheet
    for row in df.itertuples():
        text_w = set(parse_operating_hours(row.hours_text))
        sheet_w = {(row.a_first, row.a_last)}
        if row.b_first is not None:
            sheet_w.add((row.b_first, row.b_last))
        if text_w != sheet_w:
            warnings.append(f"route {row.route_no}: Route Master hours {row.hours_text!r} "
                            f"differ from Operating Windows {sorted(sheet_w)}")
        if (row.first_departure, row.last_departure) != (row.first_recomputed, row.last_recomputed) \
                or not math.isclose(row.span_hours, row.span_recomputed, abs_tol=1e-6):
            warnings.append(f"route {row.route_no}: cached first/last/span differ from typed time cells")

    df["is_loop"] = [(_place_key(o) == _place_key(d)) for o, d in zip(df.origin, df.destination)]
    loop_sheet = df.dir_a.astype(str).str.startswith("Loop")
    for rn in df.route_no[df.is_loop != loop_sheet]:
        warnings.append(f"route {rn}: is_loop (origin==destination) disagrees with Operating Windows label")

    df["headway_min_lo"] = df.freq_min_col
    df["headway_min_hi"] = df.freq_max_col
    df["f_max"] = df.headway_min_lo.map(P.headway_min_to_freq_per_h)  # shortest headway -> highest freq
    df["f_min"] = df.headway_min_hi.map(P.headway_min_to_freq_per_h)  # LONGEST headway -> lowest freq
    df["f_mid"] = (df.f_min + df.f_max) / 2.0
    fmt = lambda t: t.strftime("%H:%M")  # noqa: E731
    df["first_departure"] = df.first_departure.map(fmt)
    df["last_departure"] = df.last_departure.map(fmt)
    df["overnight"] = df.last_departure < df.first_departure

    cols = ["route_no", "origin", "destination", "is_loop", "first_departure", "last_departure",
            "span_hours", "headway_min_lo", "headway_min_hi", "f_min", "f_max", "f_mid",
            "d_r_km", "x_base", "n_via_stops",
            "service_type", "overnight", "active_am", "active_midday", "active_pm"]
    out = df[cols].copy()

    _require((out.f_min <= out.f_max).all(), "f_min > f_max somewhere")
    _require(((out.span_hours > 0) & (out.span_hours <= 24)).all(), "span outside (0, 24] h")
    _require(out[["active_am", "active_midday", "active_pm"]].isin([0, 1]).all().all(),
             "active flags must be 0/1 (cached values missing?)")
    return out, warnings


def check_model_variables(routes: pd.DataFrame, mv: pd.DataFrame, v_kmh=P.EXCEL_SPEED_KMH,
                          tol=1e-9) -> pd.DataFrame:
    """Compare the workbook's Model Variables against our recomputation. Returns one row per mismatch.

    tau = d / v [h] (the workbook's own definition, no layover); f_mid = (f_min+f_max)/2 [dep/h];
    max feasible frequency = x / tau [dep/h].
    """
    m = routes.merge(mv, on="route_no", how="outer", indicator=True, validate="one_to_one")
    issues = []
    for rn in m.route_no[m._merge != "both"]:
        issues.append(dict(route_no=rn, field="presence", excel=None, ours=None, note="route in only one sheet"))
    m = m[m._merge == "both"]
    tau = m.d_r_km / v_kmh
    checks = {
        "d_r_km": (m.xl_d, m.d_r_km),
        "x_base": (m.xl_x, m.x_base),
        "f_min": (m.xl_fmin, m.f_min),
        "f_max": (m.xl_fmax, m.f_max),
        "tau_h (d/22)": (m.xl_tau, tau),
        "f_mid": (m.xl_fmid, m.f_mid),
        "x/tau": (m.xl_xtau, m.x_base / tau),
    }
    for field, (xl, ours) in checks.items():
        for rn, a, b in zip(m.route_no, xl, ours):
            if a is None or not isinstance(a, (int, float)):
                issues.append(dict(route_no=rn, field=field, excel=a, ours=b, note="no numeric cached value"))
            elif not math.isclose(float(a), float(b), rel_tol=tol, abs_tol=tol):
                issues.append(dict(route_no=rn, field=field, excel=a, ours=b, note="value mismatch"))
    # The demand/capacity/overcrowding columns must be empty placeholders (no ridership data exists).
    for rn, D, C, o in zip(m.route_no, m.xl_D, m.xl_C, m.xl_o):
        if any(isinstance(x, (int, float)) for x in (D, C, o)):
            issues.append(dict(route_no=rn, field="D/C/o", excel=(D, C, o), ours=None,
                               note="unexpected numeric demand data"))
    return pd.DataFrame(issues, columns=["route_no", "field", "excel", "ours", "note"])


def build_stop_network(stops: pd.DataFrame, routes: pd.DataFrame):
    """Stop Coverage -> (route_stops long table, per-route features, list of warnings).

    Features per route:
        n_stops              number of stops listed for the route in Stop Coverage
        centrality_sum       sum over its stops of 'routes serving' (includes the route itself)
        centrality_mean      centrality_sum / n_stops
        n_hub_stops          stops served by >= HUB_MIN_ROUTES routes
    """
    warnings = []
    _require(len(stops) == P.N_STOPS, f"Stop Coverage has {len(stops)} stops, expected {P.N_STOPS}")
    route_set = set(routes.route_no)
    xb = routes.set_index("route_no").x_base

    long = stops.explode("routes").rename(columns={"routes": "route_no"})
    long = long.dropna(subset=["route_no"])
    unknown = sorted(set(long.route_no) - route_set)
    _require(not unknown, f"Stop Coverage mentions unknown routes: {unknown}")

    dup = long.duplicated(["stop", "route_no"])
    if dup.any():
        warnings.append(f"{int(dup.sum())} duplicated (stop, route) pairs dropped")
        long = long[~dup]
    for row in stops.itertuples():
        if len(set(row.routes)) != row.n_routes_serving:
            warnings.append(f"stop {row.stop!r}: lists {len(set(row.routes))} routes, says {row.n_routes_serving}")
        bus = int(sum(xb[r] for r in set(row.routes)))
        if bus != row.buses_passing:
            warnings.append(f"stop {row.stop!r}: sum of x_base over its routes {bus} != buses passing {row.buses_passing}")

    long = long[["route_no", "stop", "n_routes_serving"]].reset_index(drop=True)
    g = long.groupby("route_no")
    feat = pd.DataFrame({
        "n_stops": g.size(),
        "centrality_sum": g.n_routes_serving.sum(),
        "n_hub_stops": g.n_routes_serving.apply(lambda s: int((s >= P.HUB_MIN_ROUTES).sum())),
    })
    feat["centrality_mean"] = feat.centrality_sum / feat.n_stops
    feat = feat.reindex(routes.route_no)
    missing = feat.index[feat.n_stops.isna()].tolist()
    _require(not missing, f"routes with no stops in Stop Coverage: {missing}")
    feat = feat.reset_index()
    feat["n_via_stops"] = routes.n_via_stops.values
    feat["n_stops_minus_via"] = feat.n_stops - feat.n_via_stops
    feat[["n_stops", "centrality_sum", "n_hub_stops"]] = feat[["n_stops", "centrality_sum", "n_hub_stops"]].astype(int)
    return long, feat[["route_no", "n_stops", "n_via_stops", "n_stops_minus_via",
                       "centrality_sum", "centrality_mean", "n_hub_stops"]], warnings


def load_all(path=P.RAW_XLSX) -> dict:
    """Read + validate everything. Raises DataError on any hard failure. Returns a dict of frames."""
    wb_f, wb_v = _open(path)
    rm = read_route_master(wb_v)
    ow = read_operating_windows(wb_f, wb_v)
    routes, warn_r = build_routes(rm, ow)

    pdf = read_pdf_verification(wb_v, rm)
    n_ok = int((pdf.match == "OK").sum())
    _require(len(pdf) == P.N_ROUTES and n_ok == P.N_ROUTES,
             f"PDF Verification: {n_ok} OK of {len(pdf)} rows, expected {P.N_ROUTES} OK")

    mv = read_model_variables(wb_v)
    _require(len(mv) == P.N_ROUTES, f"Model Variables has {len(mv)} route rows, expected {P.N_ROUTES}")
    mv_issues = check_model_variables(routes, mv)

    stops = read_stop_coverage(wb_v)
    route_stops, features, warn_s = build_stop_network(stops, routes)
    periods = read_period_boundaries(wb_v)
    return dict(routes=routes, ow=ow, pdf=pdf, mv=mv, mv_issues=mv_issues, stops=stops,
                route_stops=route_stops, features=features, periods=periods,
                sheetnames=list(wb_v.sheetnames), wb_v=wb_v,
                warnings=warn_r + warn_s, ow_from_cache=bool(ow.from_cache.all()))


def write_processed(d: dict) -> None:
    """Write the processed CSVs. Read them back with read_routes_csv (dtype=str) so "1" stays a string."""
    P.PROCESSED.mkdir(parents=True, exist_ok=True)
    d["routes"].to_csv(P.PROCESSED / "routes.csv", index=False)
    st = d["stops"].copy()
    st["routes"] = st.routes.map(lambda xs: ", ".join(xs))
    st.to_csv(P.PROCESSED / "stops.csv", index=False)
    d["route_stops"].to_csv(P.PROCESSED / "route_stops.csv", index=False)
    d["features"].to_csv(P.PROCESSED / "route_features.csv", index=False)


def read_routes_csv(path=None) -> pd.DataFrame:
    """Read routes.csv back with route_no forced to string (pandas would turn "1" into 1)."""
    return pd.read_csv(path or P.PROCESSED / "routes.csv", dtype={"route_no": str})
